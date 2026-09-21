"""Stratified Dataset Sampler — Build sampled_dataset_1120.csv

Takes rows from 6 domain Excel files in E:\\Attack\\full data and produces a
single stratified CSV in E:\\Attack\\Attack for use in the attack pipeline.

Distribution (strictly as documented in Project_Flow.md):
  Crime_public_safety  : 175 rows (all available — this is the limiting domain)
  Celebrity_News       : 189 rows
  Disaster_BreakingNews: 189 rows
  Government_Schemes   : 189 rows
  Health_Medicine      : 189 rows
  Politics_and_Election: 189 rows
  ─────────────────────────────
  Total                : 1,120 rows

Within each domain, rows are selected using stratified sampling:
  - Aims for equal SUP / REF / NEI balance.
  - Falls back to best-effort proportional balance if a label class has
    fewer rows than the ideal share.

Label normalisation (raw → standard):
  SUPPORTED            → SUP
  REFUTED              → REF
  NOT SUPPORTED        → REF
  NOT ENOUGH INFORMATION → NEI
  SUP / REF / NEI      → kept as-is

Usage:
    python make_sampled_dataset.py
    python make_sampled_dataset.py --seed 42 --output my_sample.csv
"""

from __future__ import annotations

import argparse
import csv
import os
import random
import sys

# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------

FULL_DATA_DIR = r"E:\Attack\full data"
DEFAULT_OUTPUT = r"E:\Attack\Attack\sampled_dataset_1120.csv"
DEFAULT_SEED = 42

# Domain name → (filename, target_rows)
DOMAIN_CONFIG = {
    "Crime_PublicSafety":    ("Crime_public_safety.xlsx",    175),
    "Celebrity_News":        ("Celebrity_News.xlsx",          189),
    "Disaster_BreakingNews": ("Disaster_BreakingNews.xlsx",   189),
    "Government_Schemes":    ("Government_Schemes.xlsm",      189),
    "Health_Medicine":       ("Health_Medicine.xlsx",         189),
    "Politics_Election":     ("Politics_and_Election.xlsx",   189),
}

LABEL_MAP = {
    "SUPPORTED":              "SUP",
    "SUP":                    "SUP",
    "REFUTED":                "REF",
    "REF":                    "REF",
    "NOT SUPPORTED":          "REF",
    "NOT ENOUGH INFORMATION": "NEI",
    "NEI":                    "NEI",
}

OUTPUT_FIELDS = ["row_id", "claim", "evidence", "label", "domain", "language"]


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def load_domain(path: str, domain_name: str) -> list[dict]:
    """Load an XLSX/XLSM file and return normalised rows."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header = None
    raw_rows = []
    for row in ws.iter_rows(values_only=True):
        if header is None:
            header = [str(c).strip() if c is not None else "" for c in row]
            continue
        cell = {header[j]: row[j] for j in range(min(len(header), len(row)))}
        raw_rows.append(cell)
    wb.close()

    # Field aliases
    def _get(row, *aliases):
        for a in aliases:
            v = row.get(a)
            if v is not None:
                s = str(v).replace("\xa0", " ").strip()
                if s:
                    return s
        return ""

    out = []
    skipped = 0
    for raw in raw_rows:
        claim    = _get(raw, "claim", "claim_text", "Claim", "claim1")
        evidence = _get(raw, "evidence", "evidence_text", "Evidence")
        raw_lbl  = _get(raw, "label", "gold_label", "Label", "goldLabel").upper()
        label    = LABEL_MAP.get(raw_lbl, "")
        lang     = _get(raw, "language", "lang", "Language") or "hi"

        if not claim:
            skipped += 1
            continue
        if not label:
            # unknown label — treat as skipped
            skipped += 1
            continue

        out.append({
            "claim":    claim,
            "evidence": evidence,
            "label":    label,
            "domain":   domain_name,
            "language": lang,
        })

    if skipped:
        print(f"  [warn] {domain_name}: {skipped} rows skipped (empty claim or unknown label)")
    return out


def stratified_sample(rows: list[dict], n: int, rng: random.Random) -> list[dict]:
    """Return n rows from rows with max-balanced SUP/REF/NEI distribution.

    Algorithm:
      1. Split rows into three label buckets.
      2. Ideal share = n // 3 per bucket; remainder goes to the largest bucket.
      3. If a bucket has fewer rows than its ideal share, take all of them
         and redistribute the deficit to remaining buckets (descending size).
      4. Shuffle each bucket before slicing so selection is random.
    """
    buckets: dict[str, list[dict]] = {"SUP": [], "REF": [], "NEI": []}
    for r in rows:
        buckets.get(r["label"], buckets.setdefault(r["label"], []))
        buckets[r["label"]].append(r)

    # Shuffle each bucket
    for b in buckets.values():
        rng.shuffle(b)

    base, rem = divmod(n, 3)
    # Ideal allocations — give remainder to largest bucket
    labels_by_size = sorted(buckets.keys(), key=lambda l: len(buckets[l]), reverse=True)
    allocs = {l: base for l in buckets}
    allocs[labels_by_size[0]] += rem

    selected = []
    deficit = 0
    deferred = []  # labels that have surplus after filling their alloc

    for label in labels_by_size:
        alloc = allocs[label]
        available = len(buckets[label])
        if available >= alloc:
            selected.extend(buckets[label][:alloc])
            deferred.append((label, available - alloc))
        else:
            selected.extend(buckets[label])
            deficit += alloc - available
            print(f"    >> {label} has only {available} rows (wanted {alloc}); deficit={deficit}")

    # Fill deficit from surplus buckets (largest surplus first)
    deferred.sort(key=lambda x: x[1], reverse=True)
    for label, surplus in deferred:
        if deficit <= 0:
            break
        take = min(deficit, surplus)
        already = allocs[label] if len(buckets[label]) >= allocs[label] else len(buckets[label])
        selected.extend(buckets[label][already: already + take])
        deficit -= take

    rng.shuffle(selected)
    return selected[:n]


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def parse_args():
    p = argparse.ArgumentParser(description="Build sampled_dataset_1120.csv")
    p.add_argument("--seed",   type=int, default=DEFAULT_SEED, help="Random seed")
    p.add_argument("--output", default=DEFAULT_OUTPUT,         help="Output CSV path")
    p.add_argument("--data-dir", default=FULL_DATA_DIR,        help="Directory with domain xlsx files")
    return p.parse_args()


def main():
    args = parse_args()
    rng  = random.Random(args.seed)

    print(f"\n{'='*60}")
    print(f" Stratified Dataset Sampler — seed={args.seed}")
    print(f"{'='*60}")
    print(f" Source dir : {args.data_dir}")
    print(f" Output     : {args.output}")
    print(f"{'='*60}\n")

    all_rows: list[dict] = []
    domain_stats: dict[str, dict] = {}

    for domain_name, (filename, target) in DOMAIN_CONFIG.items():
        path = os.path.join(args.data_dir, filename)
        if not os.path.exists(path):
            print(f"[ERROR] File not found: {path}")
            sys.exit(1)

        print(f"[{domain_name}]  target={target}  file={filename}")
        rows = load_domain(path, domain_name)

        # Label breakdown before sampling
        lbl_counts = {l: sum(1 for r in rows if r["label"] == l) for l in ("SUP", "REF", "NEI")}
        print(f"  Available: {len(rows)} rows | SUP={lbl_counts['SUP']} REF={lbl_counts['REF']} NEI={lbl_counts['NEI']}")

        if len(rows) < target:
            print(f"  [WARN] Only {len(rows)} rows available but target is {target}. Taking all {len(rows)} rows.")
            target = len(rows)

        sample = stratified_sample(rows, target, rng)

        # Label breakdown after sampling
        s_counts = {l: sum(1 for r in sample if r["label"] == l) for l in ("SUP", "REF", "NEI")}
        print(f"  Sampled  : {len(sample)} rows | SUP={s_counts['SUP']} REF={s_counts['REF']} NEI={s_counts['NEI']}\n")

        domain_stats[domain_name] = {"total": len(sample), **s_counts}
        all_rows.extend(sample)

    # Assign global row_id
    rng.shuffle(all_rows)
    for i, row in enumerate(all_rows):
        row["row_id"] = i

    # Write output CSV
    os.makedirs(os.path.dirname(args.output), exist_ok=True)
    with open(args.output, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=OUTPUT_FIELDS)
        writer.writeheader()
        writer.writerows(all_rows)

    # Final summary
    total = len(all_rows)
    tot_sup = sum(1 for r in all_rows if r["label"] == "SUP")
    tot_ref = sum(1 for r in all_rows if r["label"] == "REF")
    tot_nei = sum(1 for r in all_rows if r["label"] == "NEI")

    print(f"{'='*60}")
    print(f" FINAL DATASET SUMMARY")
    print(f"{'='*60}")
    print(f" {'Domain':<25} {'Total':>6} {'SUP':>6} {'REF':>6} {'NEI':>6}")
    print(f" {'-'*51}")
    for domain, s in domain_stats.items():
        print(f" {domain:<25} {s['total']:>6} {s['SUP']:>6} {s['REF']:>6} {s['NEI']:>6}")
    print(f" {'-'*51}")
    print(f" {'TOTAL':<25} {total:>6} {tot_sup:>6} {tot_ref:>6} {tot_nei:>6}")
    print(f"{'='*60}")
    print(f"\nDone! Saved {total} rows -> {args.output}\n")


if __name__ == "__main__":
    main()
